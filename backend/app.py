# backend/app.py
from flask import Flask, request, jsonify
from flask_cors import CORS
import os
import uuid
import fitz  # PyMuPDF
import docx
import csv
import openpyxl
from dotenv import load_dotenv
import google.generativeai as genai
import pytesseract
from PIL import Image
import io

from prompts import build_prompt, build_context, parse_llm_json

# ------------------ Config ------------------

load_dotenv()
genai.configure(api_key=os.getenv("GOOGLE_API_KEY"))

DEFAULT_MODEL = os.getenv("MODEL_NAME", "gemini-1.5-flash")
DEFAULT_TEMPERATURE = float(os.getenv("TEMPERATURE", "0.2"))
MAX_CONTEXT_CHARS = int(os.getenv("MAX_CHARS", "12000"))

app = Flask(__name__)
CORS(app)

UPLOAD_FOLDER = os.path.join(os.getcwd(), 'uploads')
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

DOC_TEXTS = {}   # doc_id -> extracted text
DOC_META  = {}   # doc_id -> metadata
SESSION_MEMORY = {}  # session_id -> list of Q/A turns

# ------------------ Extraction helpers ------------------

def extract_text(filepath):
    ext = os.path.splitext(filepath)[1].lower()
    if ext == '.pdf':
        return extract_pdf(filepath)
    elif ext == '.txt':
        return extract_txt(filepath)
    elif ext == '.docx':
        return extract_docx(filepath)
    elif ext == '.csv':
        return extract_csv(filepath)
    elif ext == '.xlsx':
        return extract_xlsx(filepath)
    elif ext in ('.jpg', '.jpeg', '.png'):
        return extract_image(filepath)
    else:
        raise ValueError("Unsupported file type")

def extract_pdf(filepath):
    """Extract text from PDFs. If empty, fallback to OCR."""
    text = ""
    pages = 0
    with fitz.open(filepath) as doc:
        pages = len(doc)
        for page in doc:
            page_text = page.get_text()
            if page_text.strip():
                text += page_text
            else:
                # OCR fallback for scanned page
                pix = page.get_pixmap()
                img = Image.open(io.BytesIO(pix.tobytes("png")))
                ocr_text = pytesseract.image_to_string(img)
                text += ocr_text
    return text, pages

def extract_txt(filepath):
    with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
        return f.read(), None

def extract_docx(filepath):
    d = docx.Document(filepath)
    text = "\n".join([p.text for p in d.paragraphs])
    return text, None

def extract_csv(filepath):
    lines = []
    with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
        reader = csv.reader(f)
        for row in reader:
            lines.append(", ".join(row))
    return "\n".join(lines), None

def extract_xlsx(filepath):
    wb = openpyxl.load_workbook(filepath, data_only=True)
    text = ""
    for sheet in wb:
        for row in sheet.iter_rows(values_only=True):
            text += ", ".join([str(cell) for cell in row if cell is not None]) + "\n"
    return text, None

def extract_image(filepath):
    """OCR for image files (jpg/png)."""
    img = Image.open(filepath)
    text = pytesseract.image_to_string(img)
    return text, None

# ------------------ Routes ------------------

@app.route('/upload', methods=['POST'])
def upload_file():
    if 'file' not in request.files:
        return jsonify({"error": "No file part"}), 400

    file = request.files['file']
    if file.filename == '':
        return jsonify({"error": "No selected file"}), 400

    filename = os.path.basename(file.filename)
    filepath = os.path.join(UPLOAD_FOLDER, filename)
    file.save(filepath)

    try:
        extracted, pages = extract_text(filepath)
        if not extracted.strip():
            return jsonify({"message": "File uploaded, but no text extracted.", "text": ""}), 200

        doc_id = str(uuid.uuid4())
        DOC_TEXTS[doc_id] = extracted
        size = os.path.getsize(filepath)
        DOC_META[doc_id] = {
            "filename": filename,
            "pages": pages,
            "size_bytes": size,
        }

        return jsonify({
            "message": "File uploaded and text extracted successfully.",
            "doc_id": doc_id,
            "meta": DOC_META[doc_id],
            "text": extracted
        }), 200

    except Exception as e:
        return jsonify({"error": f"Failed to extract text: {str(e)}"}), 500

@app.route('/ask', methods=['POST'])
def ask_question():
    data = request.get_json() or {}
    question = data.get('question', '').strip()
    doc_id   = data.get('doc_id')
    session_id = data.get('session_id', 'default')
    style    = data.get('style', 'concise')
    citation_mode = bool(data.get('citation_mode', True))
    model_name = data.get('model', DEFAULT_MODEL)
    temperature = float(data.get('temperature', DEFAULT_TEMPERATURE))

    if not question:
        return jsonify({"error": "Missing 'question'"}), 400

    context_text = DOC_TEXTS.get(doc_id) if doc_id in DOC_TEXTS else next(reversed(DOC_TEXTS.values()), None)
    if not context_text:
        return jsonify({"error": "No document uploaded yet."}), 400

    try:
        # Maintain session memory
        history = SESSION_MEMORY.setdefault(session_id, [])
        history_text = "\n".join([f"Q: {h['q']}\nA: {h['a']}" for h in history])

        context = build_context(context_text + "\n" + history_text, candidate_chunks=None, max_chars=MAX_CONTEXT_CHARS)
        prompt  = build_prompt(context=context, question=question, style=style, citation_mode=citation_mode)

        model = genai.GenerativeModel(model_name)
        response = model.generate_content(prompt, generation_config={"temperature": temperature})
        raw_text = getattr(response, "text", "") or ""

        parsed = parse_llm_json(raw_text)
        parsed["doc_id"] = doc_id

        # Save into session memory
        history.append({"q": question, "a": parsed.get("answer", "")})

        return jsonify(parsed), 200
    except Exception as e:
        return jsonify({"error": f"Failed to get response: {str(e)}"}), 500

@app.get("/docs")
def list_docs():
    out = []
    for did, meta in DOC_META.items():
        out.append({"doc_id": did, **meta})
    return jsonify({"docs": out})

@app.post("/reset")
def reset_state():
    DOC_TEXTS.clear()
    DOC_META.clear()
    SESSION_MEMORY.clear()
    return jsonify({"message": "All state cleared"}), 200

@app.get("/healthz")
def healthz():
    return jsonify({"status": "ok", "model": DEFAULT_MODEL})

if __name__ == '__main__':
    app.run(port=5001)
